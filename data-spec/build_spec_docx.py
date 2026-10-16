"""Ceres database hand-over: the data dictionary (data/data-dictionary.md) as a Word document, with the identity
screen's wording (identity-screen.md) after the identity rules.   /tmp/docxenv/bin/python build_spec_docx.py"""
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).parent
DICT = HERE.parent / 'data' / 'data-dictionary.md'
IDENT = HERE / 'identity-screen.md'
OUT = HERE / 'Ceres database hand-over.docx'

INK = RGBColor(0x1A, 0x1A, 0x1A)
BRONZE = RGBColor(0x8A, 0x61, 0x16)
MUTED = RGBColor(0x5F, 0x5A, 0x52)
MONO = 'Consolas'
SANS = 'Calibri'


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_fill)
    tcPr.append(shd)


def para_shade(p, hex_fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_fill)
    pPr.append(shd)


def runs(p, text, size=None, color=None, bold=False):
    """Inline markdown: **bold** and `code`."""
    for tok in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**'):
            r = p.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith('`') and tok.endswith('`'):
            r = p.add_run(tok[1:-1]); r.font.name = MONO
            r._element.rPr.rFonts.set(qn('w:eastAsia'), MONO)
            if size: r.font.size = Pt(size - 0.5)
        else:
            r = p.add_run(tok); r.bold = bold or None
        if size and not (tok.startswith('`') and tok.endswith('`')):
            r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color


def split_row(line):
    """Split a pipe-table row on | that are not escaped and not inside `code`."""
    cells, cur, code = [], '', False
    s = line.strip()
    if s.startswith('|'): s = s[1:]
    if s.endswith('|') and not s.endswith('\\|'): s = s[:-1]
    i = 0
    while i < len(s):
        c = s[i]
        if c == '\\' and i + 1 < len(s) and s[i + 1] == '|':
            cur += '|'; i += 2; continue
        if c == '`': code = not code
        if c == '|' and not code:
            cells.append(cur.strip()); cur = ''
        else:
            cur += c
        i += 1
    cells.append(cur.strip())
    return cells


def add_table(doc, rows, width_cm):
    head, body = rows[0], rows[2:] if len(rows) > 1 and set(rows[1].replace('|', '').strip()) <= set('-: ') else rows[1:]
    hdr = split_row(head)
    data = [split_row(r) for r in body]
    n = len(hdr)
    t = doc.add_table(rows=1 + len(data), cols=n)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    # column widths in proportion to their text, within bounds
    lens = [max([len(hdr[j])] + [len(r[j]) if j < len(r) else 0 for r in data]) for j in range(n)]
    lens = [min(max(L, 6), 90) for L in lens]
    tot = sum(lens)
    widths = [Cm(width_cm * L / tot) for L in lens]
    size = 8 if n >= 5 else 8.5
    for j, h in enumerate(hdr):
        c = t.rows[0].cells[j]; c.text = ''
        runs(c.paragraphs[0], h, size=size, bold=True)
        shade(c, 'EFEBE4')
    for i, r in enumerate(data):
        for j in range(n):
            c = t.rows[i + 1].cells[j]; c.text = ''
            runs(c.paragraphs[0], r[j] if j < len(r) else '', size=size)
    for row in t.rows:
        for j, c in enumerate(row.cells):
            c.width = widths[j]
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(0)
    # repeat the header row on each page
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement('w:tblHeader'); th.set(qn('w:val'), 'true'); trPr.append(th)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def render(doc, md_lines, width_cm):
    i, para = 0, []

    def flush():
        if para:
            p = doc.add_paragraph(); runs(p, ' '.join(x.strip() for x in para), size=10)
            para.clear()

    while i < len(md_lines):
        line = md_lines[i]
        s = line.rstrip('\n')
        if s.startswith('```'):
            flush(); i += 1
            while i < len(md_lines) and not md_lines[i].startswith('```'):
                p = doc.add_paragraph(); r = p.add_run(md_lines[i].rstrip('\n')); r.font.name = MONO; r.font.size = Pt(8.5)
                p.paragraph_format.space_after = Pt(0); para_shade(p, 'F4F2EE')
                i += 1
            i += 1; doc.add_paragraph().paragraph_format.space_after = Pt(2); continue
        m = re.match(r'^(#{1,3}) (.*)$', s)
        if m:
            flush()
            level = len(m.group(1))
            if level == 1:
                i += 1; continue          # the document title is written separately
            h = doc.add_heading(level=level - 1)
            runs(h, m.group(2), size={2: 15, 3: 12}[level], color=BRONZE if level == 2 else INK)
            i += 1; continue
        if s.startswith('|'):
            flush(); rows = []
            while i < len(md_lines) and md_lines[i].startswith('|'):
                rows.append(md_lines[i].rstrip('\n')); i += 1
            add_table(doc, rows, width_cm); continue
        m = re.match(r'^(\s*)(- |\d+\. )(.*)$', s)
        if m:
            flush()
            numbered = m.group(2)[0].isdigit()
            text = (m.group(2) if numbered else '') + m.group(3)
            i += 1
            while i < len(md_lines) and md_lines[i].strip() and re.match(r'^\s{2,}\S', md_lines[i]) and not re.match(r'^\s*(- |\d+\. )', md_lines[i]):
                text += ' ' + md_lines[i].strip(); i += 1
            p = doc.add_paragraph(style='List Bullet' if not numbered else None)
            if numbered:
                p.paragraph_format.left_indent = Cm(0.6); p.paragraph_format.first_line_indent = Cm(-0.6)
            if len(m.group(1)) >= 2 and not numbered:
                p.paragraph_format.left_indent = Cm(1.2)
            runs(p, text, size=10)
            p.paragraph_format.space_after = Pt(2)
            continue
        if not s.strip():
            flush(); i += 1; continue
        if re.match(r'^\s*<!--.*-->\s*$', s):
            flush(); i += 1; continue          # the generator's markers are not for readers
        para.append(s); i += 1
    flush()


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for side in ('left_margin', 'right_margin'): setattr(sec, side, Cm(1.8))
    sec.top_margin = sec.bottom_margin = Cm(1.6)
    width_cm = 29.7 - 3.6
    st = doc.styles['Normal']; st.font.name = SANS; st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), SANS)
    for name in ('Heading 1', 'Heading 2'):
        hs = doc.styles[name]; hs.font.name = SANS; hs.element.rPr.rFonts.set(qn('w:eastAsia'), SANS)
        hs.font.color.rgb = INK
    # proofing language
    for s in (doc.styles['Normal'],):
        rPr = s.element.get_or_add_rPr(); lang = OxmlElement('w:lang'); lang.set(qn('w:val'), 'en-GB'); rPr.append(lang)

    t = doc.add_paragraph(); r = t.add_run('C E R E S'); r.font.size = Pt(26); r.font.name = 'Georgia'
    sub = doc.add_paragraph(); runs(sub, 'DATABASE HAND-OVER · QUESTIONNAIRE V3', size=10, color=BRONZE)
    d = doc.add_paragraph(); runs(d, '16 October 2026', size=10, color=MUTED)

    md = DICT.read_text().splitlines()
    ident = IDENT.read_text().splitlines() if IDENT.exists() else []
    # put the identity screen's wording straight after the identity rules (before "## Tables and columns")
    k = next(j for j, l in enumerate(md) if l.startswith('## Tables and columns'))
    md = md[:k] + ident + [''] + md[k:]
    render(doc, md, width_cm)

    footer = sec.footer.paragraphs[0]; runs(footer, 'Ceres · database hand-over · questionnaire v3', size=8, color=MUTED)
    doc.core_properties.title = 'Ceres: database hand-over (questionnaire v3)'
    doc.core_properties.author = 'Haresh'
    doc.save(OUT)
    print(OUT)


if __name__ == '__main__':
    main()
